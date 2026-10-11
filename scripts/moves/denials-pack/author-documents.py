"""Offline document authoring. Use the bundled Python runtime; no data-plane calls."""
import json
import sys
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor, white
from reportlab.lib.utils import simpleSplit

out = Path(sys.argv[1])
s = json.loads((out / 'scenario.json').read_text())
t = s['trailingYear']
a = s['assumptions']
notice = s['notice']
fmt = lambda v: f'{v:,.2f}'
usd = lambda v: f'${v:,.2f}'

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.75)
sec.header_distance = sec.footer_distance = Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(13)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
sec.header.paragraphs[0].text=notice
sec.header.paragraphs[0].style='Caption'
foot=sec.footer.paragraphs[0]
foot.text='Synthetic discovery brief  |  '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
doc.core_properties.title='Claims Denial Prevention Intelligence'
doc.core_properties.author='AbarVa synthetic fixture'
doc.core_properties.created=datetime(2026,10,10,tzinfo=timezone.utc)
def p(text):doc.add_paragraph(text)
def h(text):doc.add_heading(text,level=1)
def table(headers,rows,widths):
    tbl=doc.add_table(rows=1,cols=len(headers));tbl.autofit=False
    for i,v in enumerate(headers):tbl.rows[0].cells[i].text=v
    for row in rows:
        cells=tbl.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for r,row in enumerate(tbl.rows):
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]);pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17283C' if r==0 else ('F1F4F7' if r%2 else 'FFFFFF'));pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                for run in para.runs:
                    run.font.size=Pt(9)
                    if r==0:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    header=OxmlElement('w:tblHeader');tbl.rows[0]._tr.get_or_add_trPr().append(header)
    return tbl

doc.add_paragraph('Claims Denial Prevention Intelligence','Title')
p(notice)
doc.add_paragraph('Discovery request for the invented revenue cycle','Subtitle')
p('The proposed sponsor role is Revenue cycle director. The ask is to test whether reviewed administrative checks before submission can prevent avoidable denials, route corrections to the source team and trace subsequent collected cash. This brief and every source are invented. No source represents a real institution, payer policy, patient or Finance attestation.')
h('Why open this Move')
p('The synthetic revenue-cycle team receives denial feedback after submission and works a manually assembled queue. Source teams cannot reliably distinguish expired coverage, authorization mismatch, administrative coding errors and missing attachments. Corrected submissions can inflate the denominator if claim versions are not joined back to the original claim.')
table(['Trailing year through 2026 Q2','Synthetic baseline'],[
    ['Original submitted claims',f"{t['claims']:,}"],
    ['First initial denials',f"{t['denials']:,}"],
    ['Initial denial rate',f"{t['denialRate']:.4%}"],
    ['Gross charges',usd(t['grossCharges'])],
    ['Denied allowed amount',usd(t['deniedAllowed'])],
    ['Recovered cash on denied claims',usd(t['recoveries'])],
    ['Final write offs',usd(t['writeOffs'])],
    ['Rework hours',fmt(t['reworkHours'])],
],[3.5,3.3])
p('Basis: eight-quarter submission cohorts from 2024 Q3 through 2026 Q2. The trailing year is the final four quarters. Original submissions count once; only the first denial counts. Every invented cohort reaches final disposition within 90 days and is mature by 10 October 2026. Recovered cash, write-offs and contractual adjustments partition denied allowed amounts. Gross charges are not a cash benefit basis.')
h('Scope and boundaries')
p('Include all four invented facilities, all four invented payer cohorts and all three administrative service groups. Assess original claim identity, service-date coverage, authorization scope, reviewed reason codes, source-team correction routing and receipt reconciliation. Exclude clinical decision support, autonomous submission, real PHI, payer contract changes and source-platform replacement.')

doc.add_page_break()
h('The invented operating footprint')
table(['Facility ID','Invented organisation'],[[x['id'],x['name']] for x in s['facilities']],[1.1,5.7])
p('Service groups: SL-01 Hospital ambulatory; SL-02 Hospital inpatient; SL-03 Professional services. These are synthetic administrative cohort labels, not clinical recommendations. All facility and service combinations are invented for the demonstration.')
table(['Payer ID','Invented payer','Claim share'],[[x['id'],x['name'],f"{x['share']:.0%}"] for x in s['payers']],[1.1,4.5,1.2])
h('Ownership and decision rights')
p('The Revenue cycle director owns the business outcome. Patient access and authorization leads own their administrative checks; the Coding manager and Documentation operations lead own correction rules. The Cash application manager owns receipt linkage. The Finance business partner validates monetary baselines and collection attribution. The Data platform owner owns access and lineage. These are proposed roles; the owner must assign actual authorized Move participants.')
h('Success hypotheses for human review')
p(f"Propose an initial denial-rate target of {a['targetDenialRate']:.0%} against the reviewed synthetic baseline. Propose preventing {a['preventionShare']:.0%} of final write-off exposure, with {a['attribution']:.0%} attribution and {a['probability']:.0%} probability. Propose reducing rework hours by {a['reworkReductionShare']:.0%} and days in receivables from {a['currentDaysAR']} to {a['targetDaysAR']}. These are planning hypotheses awaiting confirmed register rows and human review, not promised results.")
p('Credit retained allowed revenue only when incremental cash is collected against a matched comparator. Existing recovered receipts cannot be counted again. Administrative time remains capacity with zero monetary value until a role or contract release is recorded. Faster collection changes timing and must not create a second annual revenue lever.')
h('Discovery evidence and unanswered questions')
p('Review the cohort workbook, taxonomy, source inventory, interviews, finance baseline, current workflow and invented policy excerpts together. Test whether service-date versions and resubmission IDs join; whether source owners can resolve each primary reason; and whether receipts identify the denied claim. Real access, causal prevention effect, adoption and incremental collections remain untested.')

doc.add_page_break()
h('Options brought to the design workshop')
p('Governed batch prevention and recovery is the proposed direction for review: a versioned administrative decision ledger, reviewed source-team routing and receipt reconciliation over existing systems. Its advantage is a traceable daily queue; its limitation is batch freshness and the need to check eligibility again before submission. This is an option proposal, not a selected or approved architecture.')
p('A synchronous administrative-check service is an alternative. It could check immediately before submission but needs source-service access, contractual integration and reviewed downtime behaviour. Manual source-team checklists are the minimum-change comparator: low integration cost but weak version lineage and continued manual queue assembly. The owner must compare the options actually served in the product, mark coverage and choose the route personally.')
h('Counting the delivery scope once')
table(['Work block','Proposed components'],[
  ['Shared foundation','6 sources; 28 source tables; 9 standard entities; 0 views; 12 design rows; 24 validation rows'],
  ['UC-1 Submission risk review','0 new sources or tables; 3 entities; 4 views; 10 design rows; 20 validation rows'],
  ['UC-2 Correction routing and recovery','0 new sources or tables; 4 entities; 5 views; 12 design rows; 24 validation rows'],
],[2.2,4.6])
p('The shared foundation connects the six inventory sources once. Use-case counts cover additional semantic and queue work only, so reused source tables are not counted again. Unit-hour rows are invented review proposals in the seed rebind; delivery rates must come exclusively from the versioned pricing-engine-v1 cost foundation and the selected pod. No priced estimate or ROM approval is asserted here.')
h('Investment and measurement conditions')
p(f"The proposed budget ceiling is {usd(a['budgetCeiling'])}. Cost for the value case must come from the current owner-approved ROM; the ceiling is not another cost basis. A proposed ramp starts in month {a['benefitStartMonth']}, runs for {a['rampMonths']} months and uses a conservative {a['currentPaymentLagMonths']}-month collection lag. The current monthly engine cannot convert an eight-day collection improvement into an invented full-month cash shift.")
p('Before any handoff, the owner reviews the exact seed rows, approves the scoped rebind by name, reviews and approves the priced estimate in P3, then checks the value engine and gate. The resulting documents remain synthetic. P5 prepares handoff to external execution and Tower measurement; it does not claim that execution has started.')
doc.save(out/'01-use-case-brief.docx')

interviews=[
('Revenue cycle director','Outcome and throughput','The review meeting receives a total denial count, but the team cannot tell whether yesterday\'s correction actually removed the first cause. A new spreadsheet becomes the queue when a payer remittance batch lands. The director wants source-team ownership before submission, with an override when a rule is wrong.','We learn which check was missing after the claim has already been returned.','A reviewed primary reason and a visible owner are required before a queue is called actionable. Clean claims must not wait behind unresolved exceptions. The director proposes a shadow comparison before any operational rollout.'),
('Patient access lead','Eligibility and authorization handoff','Coverage responses arrive with versions, but the billing export retains the latest result rather than the service-date result. Authorization scope is tracked separately; an identifier can exist while the approved service does not match. Staff copy a response into a worksheet and call another team when the dates disagree.','A green coverage response today does not explain what was checked on the service date.','Use the service-date coverage version and authorization interval. Keep missing data as an exception rather than treating it as eligible. Patient access certifies eligibility rules; the authorization lead certifies authorization scope. Access must be purpose-bound and approved.'),
('Coding manager','Administrative edits and accountable routing','Reason-code labels drift between analysts. An attachment problem can be entered as a coding problem because the code description is convenient. Queue rows have no shared original-claim identity, so repeat submissions look like additional demand. The manager wants separate reviewed administrative checks and a way to reject a mistaken rule.','A reason description is not the same thing as a reviewed root cause.','Review D-CO with the administrative code pair and source version. D-DO belongs to documentation operations. The taxonomy must retain an unresolved bucket in a live implementation; this fixture contains only reviewed invented groups. No clinical coding advice or autonomous code change is proposed.'),
('Finance business partner','Cash evidence and benefit conversion',f"The trailing-year baseline contains {t['claims']:,} original claims and {t['denials']:,} first denials. Final write-offs are {usd(t['writeOffs'])}; cash recovered on denied claims is {usd(t['recoveries'])}. Rework consumes {fmt(t['reworkHours'])} hours. Gross charges of {usd(t['grossCharges'])} are not collected value. These are invented mature cohorts, not independently attested financials.",'Show me the receipt and the comparator before calling prevented loss a cash benefit.','A reduction in rework time is capacity at zero cash without a recorded release. No collection-vendor contract release is proposed. Faster AR changes receipt timing; do not add the receivables stock as annual revenue. Budget is a ceiling; the approved ROM is the investment cost.'),
('Data platform owner','Identity lineage and source access','The six invented sources use different episode, submission and receipt identifiers. Corrected claim versions must map to the original claim. Remittance reason ordering changes between batches, and one acknowledgement feed lacks a timezone. The platform owner wants restricted landing, reviewed identity mappings and a quality quarantine before products read anything.','Keep the source version beside the decision so we can reproduce the exception.','No extract or entitlement exists yet. Source-owner approval and purpose-bound access precede ingestion. Canonical records carry identity and provenance; product screens project those records. Data platform owns releases of reviewed slices, and finance owns monetary certification.'),
]
lines=[f'# Synthetic role interview notes\n\n{notice}\n\nThese five interviews are invented workshop records dated 10 October 2026. Quotes are fictional; no participant or organisation was interviewed. Roles do not give a real sign-off.\n']
for i,(role,topic,notes,quote,boundary) in enumerate(interviews,1):
    lines += [f'## Interview {i} {role}\n\n### Topic\n\n{topic}\n\n### Session notes\n\n{notes}\n\n### Invented quote\n\n> {quote}\n\n### Proposed decision and open conditions\n\n{boundary}\n']
lines += ['## Candidate causal findings for human review\n\n- Coverage and authorization versions are missing from the pre-submission decision record.\n- Inconsistent denial reason mapping leaves manual correction queues without a shared accountable routing rule.\n- Recovery receipts lack a reconciled link to the original denied claim and its final disposition.\n\nThese proposals carry no accepted status or confirmed ranking. The owner links actual approved evidence on the new Move, reviews each cause and confirms the order.\n']
(out/'05-interview-notes.md').write_text('\n'.join(lines))

navy=HexColor('#17283C'); blue=HexColor('#315E86'); grey=HexColor('#627184'); pale=HexColor('#EDF2F6')
def title(c,text,sub):
    c.setFillColor(navy);c.setFont('Helvetica-Bold',23);c.drawString(42,550,text)
    c.setFont('Helvetica',10);c.setFillColor(grey);c.drawString(42,526,notice)
    c.setFont('Helvetica',11);c.drawString(42,502,sub)
def wrap(c,text,x,y,width,font='Helvetica',size=10,leading=14,color=navy):
    c.setFont(font,size);c.setFillColor(color)
    for line in simpleSplit(text,font,size,width):c.drawString(x,y,line);y-=leading
    return y

c=canvas.Canvas(str(out/'07-current-workflow.pdf'),pagesize=(792,612),invariant=1)
c.setTitle('Synthetic current denial workflow');title(c,'Current denial workflow','Invented administrative process from submission to final cash disposition')
nodes=[('Check and submit','Patient access and billing','Latest coverage result and separate authorization tracker.'),('Receive denial','Clearinghouse and remits','First denial recorded after submission; preserve source batch.'),('Map and route','Revenue cycle analyst','Manual reason mapping routes an exception to its source team.'),('Correct and resubmit','Source team lead','Correction loops to submission; original identity must remain.'),('Reconcile disposition','Cash application manager','Receipt, final write-off or adjustment closes the cohort.')]
xs=[42,187,332,477,622];bw=128
for x,(head,role,desc) in zip(xs,nodes):
    c.setFillColor(pale);c.roundRect(x,285,bw,155,8,stroke=0,fill=1)
    wrap(c,head,x+11,414,bw-22,'Helvetica-Bold',11,14)
    wrap(c,role,x+11,378,bw-22,'Helvetica',9,12,color=blue)
    wrap(c,desc,x+11,339,bw-22,'Helvetica',10,13)
for x in xs[:-1]:
    c.setStrokeColor(blue);c.setLineWidth(1.4);c.line(x+bw+3,365,x+bw+14,365);c.line(x+bw+14,365,x+bw+10,369);c.line(x+bw+14,365,x+bw+10,361)
c.setStrokeColor(blue);c.line(541,278,541,253);c.line(541,253,106,253);c.line(106,253,106,279)
c.setFont('Helvetica',9);c.setFillColor(blue);c.drawString(223,237,'Correction loop retains original claim identity; no new denominator claim')
wrap(c,f"Invented feedback delay: {a['detectionLagDays']} days. Trailing year: {t['claims']:,} original submissions; {t['denials']:,} first denials; {fmt(t['reworkHours'])} rework hours. Days in accounts receivable: {a['currentDaysAR']}.",42,205,708,size=11,leading=16)
wrap(c,'Observed gaps in the fictional workshop: service-date versions are missing; manual reason mapping lacks a shared routing rule; recovery receipts lack a reconciled original-claim link. These are proposed causes, not accepted findings.',42,155,708,size=10,leading=14)
wrap(c,'Sources: 02-denials-baseline.xlsx, 04-source-system-inventory.xlsx and 05-interview-notes.md. The diagram describes the invented current process; the proposed future design is not deployed.',42,99,708,size=9,leading=13,color=grey)
c.setFont('Helvetica',9);c.drawString(42,37,'AbarVa synthetic discovery source  |  1 of 1');c.save()

c=canvas.Canvas(str(out/'08-payer-policy-excerpts.pdf'),pagesize=(612,792),invariant=1)
c.setTitle('Invented payer administrative policy excerpts')
for page,(payer,policies) in enumerate([(s['payers'][0],s['reasons'][:2]),(s['payers'][1],s['reasons'][2:])],1):
    c.setFillColor(navy);c.setFont('Helvetica-Bold',20);c.drawString(42,741,'Invented payer policy excerpts')
    c.setFont('Helvetica',10);c.setFillColor(grey);c.drawString(42,717,notice)
    y=wrap(c,f"Fictional document for {payer['name']} ({payer['id']}). Payer submission share in the synthetic baseline: {payer['share']:.0%}. These administrative examples are not actual payer policies or legal or clinical guidance.",42,679,528,size=11,leading=16)
    for r in policies:
        y-=27;c.setFillColor(blue);c.setFont('Helvetica-Bold',14);c.drawString(42,y,f"{r['policy']}  {r['group']}  {r['id']}");y-=25
        clauses={
          'D-EL':'The invented coverage check must retain the subscriber-match response and coverage version effective on the service date. A later response does not retrospectively prove the earlier coverage check. An absent version routes the claim to patient access for review.',
          'D-AU':'The invented authorization rule requires the reference identifier, validity interval and approved administrative service scope to match the submitted claim. An identifier alone is insufficient. A mismatch routes to the authorization team; no automated waiver is allowed.',
          'D-CO':'The invented administrative edit compares the submitted procedure and billing modifiers to a reviewed local mapping. Overrides need a recorded reviewer and a reason. An edit does not recommend clinical treatment or replace professional coding review.',
          'D-DO':'The invented attachment rule requires a document reference and a completeness check before administrative submission. A missing reference routes to documentation operations. The fixture contains no actual clinical document and requires no real patient information.',
        }
        y=wrap(c,clauses[r['id']],42,y,528,size=11,leading=17)
        y-=14;y=wrap(c,f"Proposed owner role: {r['owner']}. Taxonomy preventability hypothesis: {r['preventable']:.0%}, invented for discovery and awaiting human validation. This percentage is not a published payer performance statistic.",42,y,528,size=10,leading=15)
        y-=17;y=wrap(c,'Review test: reconstruct the service-date decision from source versions, keep exceptions unresolved if required fields are missing, and record the reviewer before treating the rule as covered.',42,y,528,size=10,leading=15)
    wrap(c,'The other invented payer cohorts follow the same reviewed taxonomy for this fixture only. No real payer correspondence, payment incentive, tier or clinical programme is represented. Source links: 03-denial-reason-taxonomy.xlsx and 04-source-system-inventory.xlsx.',42,125,528,size=9,leading=14,color=grey)
    c.setFont('Helvetica',9);c.drawString(42,43,f'AbarVa synthetic discovery source  |  {page} of 2');c.showPage()
c.save()
print(json.dumps({'docx':1,'pdf':2,'interviews':5,'notice':notice}))
